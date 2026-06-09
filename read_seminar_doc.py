from docx import Document

def read_docx(file_path):
    doc = Document(file_path)
    full_text = []
    
    # 读取段落
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    
    # 读取表格
    for table in doc.tables:
        full_text.append("\n=== 表格内容 ===")
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                row_text.append(cell.text)
            full_text.append(" | ".join(row_text))
    
    return "\n".join(full_text)

# 读取文档
file_path = r"C:\Users\jianlinw\Desktop\AI工具专题材料\科技项目双月度研讨会主题.docx"
content = read_docx(file_path)

# 保存到文件
output_file = r"C:\AI学习资料\mesheer\seminar_content.txt"
with open(output_file, 'w', encoding='utf-8') as f:
    f.write(content)

print(f"文档内容已提取到: {output_file}")
print(f"总长度: {len(content)} 字符")
print("\n===== 文档内容预览 =====")
print(content[:2000])
