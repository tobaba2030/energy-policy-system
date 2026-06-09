# -*- coding: utf-8 -*-
import os, sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

def set_table_border(table, color='333333', sz='6'):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}></w:tblPr>')
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>')
    tblPr.append(borders)

def add_hdg(doc, text, level=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level==1 else 12)
    pf.space_after = Pt(8)
    pf.line_spacing = Pt(26)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_name = 'SimHei'; fs = 22
        pf.space_before = Pt(24); pf.space_after = Pt(18)
    elif level == 1:
        font_name = 'SimHei'; fs = 16
    elif level == 2:
        font_name = 'KaiTi'; fs = 15
    elif level == 3:
        font_name = 'KaiTi'; fs = 13
    else:
        font_name = 'FangSong'; fs = 12
    run = p.add_run(text)
    run.font.size = Pt(fs)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.bold = True
    return p

def add_para(doc, text, font_name='FangSong', font_size=12, bold=False, align=None, indent=0, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = Pt(22)
    if indent:
        pf.first_line_indent = Pt(indent)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = color
    return p

def add_img(doc, img_path, width=5.5, caption=''):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(img_path):
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.font.size = Pt(9)
        cr.font.name = 'SimSun'
        cr._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

def add_br(doc):
    doc.add_page_break()

def add_table_row(table, row_data, widths, row_style='normal'):
    tr = table.add_row()
    for j, txt in enumerate(row_data):
        cell = tr.cells[j]; cell.text = ''
        p = cell.paragraphs[0]
        if j < len(widths):
            cell.width = widths[j]
        if row_style == 'header':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fn = 'SimHei'; fs = 10; bd = True
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A237E" w:val="clear"/>')
            cell._tc.get_or_add_tcPr().append(shading)
            clr = RGBColor(0xFF, 0xFF, 0xFF)
        elif row_style == 'topic':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fn = 'KaiTi'; fs = 10; bd = True; clr = None
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EAF6" w:val="clear"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        elif row_style == 'label':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fn = 'SimHei'; fs = 10; bd = True; clr = None
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EAF6" w:val="clear"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 4 else WD_ALIGN_PARAGRAPH.LEFT
            fn = 'FangSong'; fs = 9; bd = False; clr = None
        r = p.add_run(txt)
        r.font.size = Pt(fs)
        r.font.name = fn
        r._element.rPr.rFonts.set(qn('w:eastAsia'), fn)
        if bd:
            r.font.bold = True
        if clr:
            r.font.color.rgb = clr

def build():
    sys.path.insert(0, OUT_DIR)
    import content_data_v3 as C

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)

    print('Document initialized...')

    # ===== COVER =====
    add_para(doc, '', font_size=6)
    add_para(doc, C.TITLE1, font_name='SimHei', font_size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, '', font_size=10)
    add_para(doc, C.TITLE2, font_name='KaiTi', font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, '', font_size=10)
    add_para(doc, '场景名称：虚拟电厂多时空尺度智能协同运营', font_name='FangSong', font_size=14,
             bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, '', font_size=10)

    info = [
        ('申报单位', '（待填写）'),
        ('联合申报单位', '贵州电网有限责任公司、朗新科技集团股份有限公司\n'
         '贵州矿能集团有限公司、贵州国能科技有限公司\n'
         '黔能三安新能源科技有限公司'),
        ('项目负责人', '（待填写）'),
        ('联系方式', '（待填写）'),
        ('申报日期', '2026年6月'),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(t, '333333', '6')
    for i, (k, v) in enumerate(info):
        c0, c1 = t.cell(i,0), t.cell(i,1)
        c0.width = Cm(3.5); c1.width = Cm(11)
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(k); r0.font.size = Pt(12); r0.font.name = 'SimHei'
        r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei'); r0.font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EAF6" w:val="clear"/>')
        c0._tc.get_or_add_tcPr().append(shading)
        p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(v); r1.font.size = Pt(11); r1.font.name = 'FangSong'
        r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
    add_br(doc)
    print('Cover done')

    # ===== SECTION 1: 项目背景与意义 =====
    add_hdg(doc, '一、项目背景与意义', 1)

    add_hdg(doc, '（一）行业背景与政策环境', 2)
    for para_text in C.SEC1_1_1:
        add_para(doc, para_text, indent=24)
    add_img(doc, os.path.join(OUT_DIR, 'fig_vpp_arch.png'), 5.5,
            '图1-1  虚拟电厂多时空尺度智能协同运营业务架构')

    add_hdg(doc, '（二）问题分析与需求痛点', 2)
    add_para(doc, C.SEC1_1_2, indent=24)
    for title, detail in C.SEC1_1_2_LIST:
        add_para(doc, title, font_size=12, bold=True, indent=0)
        add_para(doc, detail, indent=24)

    add_hdg(doc, '（三）项目建设的必要性与意义', 2)
    for para_text in C.SEC1_2:
        add_para(doc, para_text, indent=24)

    add_br(doc)
    print('Section 1 done')

    # ===== SECTION 2: 主要研究内容 =====
    add_hdg(doc, '二、主要研究内容', 1)
    add_para(doc, C.SEC2_INTRO, indent=24)

    # Research table
    research_table = doc.add_table(rows=1, cols=4)
    research_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(research_table, '333333', '6')
    widths_r = [Cm(2.0), Cm(4.5), Cm(3.5), Cm(4.5)]

    # Header row
    hdr = research_table.rows[0]
    header_labels = ['子任务编号', '子任务名称', '牵头方', '参与方']
    for i, hd in enumerate(header_labels):
        cell = hdr.cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(hd); r.font.size = Pt(10); r.font.name = 'SimHei'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei'); r.font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A237E" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.width = widths_r[i]

    for topic, subtasks in C.SEC2_RESEARCH:
        # Topic row
        tr = research_table.add_row()
        # Merge all columns for topic
        tr.cells[0].merge(tr.cells[3])
        cell = tr.cells[0]; cell.text = ''
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(topic); r.font.size = Pt(10); r.font.name = 'KaiTi'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'KaiTi'); r.font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EAF6" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        cell.width = Cm(14.5)

        # Subtask rows
        for sub in subtasks:
            sr = research_table.add_row()
            for j, txt in enumerate(sub):
                cell = sr.cells[j]; cell.text = ''
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j >= 2 else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(txt); r.font.size = Pt(9); r.font.name = 'FangSong'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
                cell.width = widths_r[j]

    add_br(doc)
    print('Section 2 done')

    # ===== SECTION 3: 创新点 =====
    add_hdg(doc, '三、创新点', 1)
    for title, detail, compare in C.SEC3_INNOVATIONS:
        add_hdg(doc, title, 3)
        add_para(doc, detail, indent=24)
        add_para(doc, compare, font_name='FangSong', font_size=9.5, indent=24,
                 color=RGBColor(0x15, 0x65, 0xC0))
    add_img(doc, os.path.join(OUT_DIR, 'fig_vpp_sysarch.png'), 5.5,
            '图3-1  虚拟电厂多时空尺度智能协同运营系统总体架构')

    add_br(doc)
    print('Section 3 done')

    # ===== SECTION 4: 资金情况 =====
    add_hdg(doc, '四、资金情况', 1)
    add_para(doc, C.SEC4_INTRO, indent=24)

    budget_table = doc.add_table(rows=1, cols=2)
    budget_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(budget_table, '333333', '6')

    # Header
    hp = budget_table.rows[0]
    for i, hd in enumerate(['项目', '内容']):
        cell = hp.cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(hd); r.font.size = Pt(11); r.font.name = 'SimHei'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei'); r.font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A237E" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    hp.cells[0].width = Cm(4)
    hp.cells[1].width = Cm(12)

    for row_data in C.SEC4_BUDGET:
        tr = budget_table.add_row()
        for j, txt in enumerate(row_data):
            cell = tr.cells[j]; cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.font.size = Pt(10)
            fn = 'FangSong' if j == 1 else 'SimHei'
            r.font.name = fn; r._element.rPr.rFonts.set(qn('w:eastAsia'), fn)
            if j == 0:
                r.font.bold = True
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8EAF6" w:val="clear"/>')
                cell._tc.get_or_add_tcPr().append(shading)
        tr.cells[0].width = Cm(4)
        tr.cells[1].width = Cm(12)

    add_br(doc)
    print('Section 4 done')

    # ===== SECTION 5: 项目整体计划 =====
    add_hdg(doc, '五、项目整体计划（包含主要阶段）', 1)

    plan_table = doc.add_table(rows=1, cols=5)
    plan_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(plan_table, '333333', '6')
    widths_p = [Cm(1.0), Cm(3.5), Cm(2.0), Cm(1.8), Cm(6.2)]

    # Header
    ph = plan_table.rows[0]
    for i, hd in enumerate(['序号', '任务名称', '计划开始时间', '预计工期', '任务描述']):
        cell = ph.cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(hd); r.font.size = Pt(10); r.font.name = 'SimHei'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei'); r.font.bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A237E" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.width = widths_p[i]

    for row_data in C.SEC5_PLAN:
        tr = plan_table.add_row()
        for j, txt in enumerate(row_data):
            cell = tr.cells[j]; cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 4 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(txt); r.font.size = Pt(9); r.font.name = 'FangSong'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
            cell.width = widths_p[j]

    add_img(doc, os.path.join(OUT_DIR, 'fig_vpp_gantt.png'), 5.5,
            '图5-1  虚拟电厂多时空尺度智能协同运营项目实施计划甘特图（15个月）')

    add_br(doc)
    print('Section 5 done')

    # ===== SECTION 6: 待协调事项 =====
    add_hdg(doc, '六、待协调事项', 1)
    add_para(doc, '为确保项目顺利推进和实施，以下事项需在项目启动前期重点协调落实：', indent=24)
    for title, detail in C.SEC6_COORD:
        add_hdg(doc, title, 3)
        add_para(doc, detail, indent=24)

    print('Section 6 done')

    # ===== SAVE =====
    out_path = os.path.join(OUT_DIR, '虚拟电厂多时空尺度智能协同运营_申报方案_v3.docx')
    doc.save(out_path)
    print(f'SUCCESS: {out_path}')
    print(f'Size: {os.path.getsize(out_path):,} bytes')

if __name__ == '__main__':
    build()
