#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
技术报告Word文档生成器
用于将Markdown格式的技术报告转换为格式规范的Word文档
"""

import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class TechnicalReportGenerator:
    def __init__(self):
        self.doc = Document()
        self.setup_document()

    def setup_document(self):
        """设置文档默认样式"""
        # 设置默认字体
        style = self.doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)
        font.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 设置段落格式
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.first_line_indent = Inches(0.3)
        
    def add_title(self, text):
        """添加标题页"""
        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(text)
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        
        self.doc.add_paragraph()
        self.doc.add_paragraph()
        
    def add_subtitle(self, text):
        """添加副标题"""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        
    def add_heading(self, text, level=1):
        """添加章节标题"""
        para = self.doc.add_heading(text, level=level)
        if level == 1:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in para.runs:
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        elif level == 2:
            for run in para.runs:
                run.font.size = Pt(16)
                run.font.bold = True
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        elif level == 3:
            for run in para.runs:
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    def add_paragraph(self, text):
        """添加正文段落"""
        # 处理加粗文本
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.first_line_indent = Inches(0.3)
        
        # 简单处理加粗标记
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = para.add_run(part[2:-2])
                run.font.bold = True
            else:
                run = para.add_run(part)
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    def add_list_item(self, text, level=1):
        """添加列表项"""
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.3 * level)
        para.paragraph_format.line_spacing = 1.5
        
        if text.strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.')):
            # 有序列表
            para.text = text.strip()
        elif text.strip().startswith(('●', '○', '■', '□', '•', '-', '*')):
            # 无序列表
            para.text = '    ' + text.strip()
        else:
            para.text = text.strip()
            
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    def add_table(self, headers, rows):
        """添加表格"""
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        # 添加表头
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            for paragraph in hdr_cells[i].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(11)
        
        # 添加数据行
        for row in rows:
            row_cells = table.add_row().cells
            for i, cell_text in enumerate(row):
                row_cells[i].text = cell_text
                for paragraph in row_cells[i].paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
        
        self.doc.add_paragraph()

    def add_code_block(self, text):
        """添加代码块"""
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.left_indent = Inches(0.5)
        para.paragraph_format.right_indent = Inches(0.5)
        
        run = para.add_run(text)
        run.font.size = Pt(10)
        run.font.name = 'Consolas'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'F2F2F2')
        para._element.get_or_add_pPr().insert_element_before(shading, 'w:spacing')
        
    def add_figure_placeholder(self, caption):
        """添加图表占位符"""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"[{caption}]")
        run.font.size = Pt(11)
        run.font.italic = True
        self.doc.add_paragraph()
        
    def add_page_break(self):
        """添加分页符"""
        self.doc.add_page_break()

    def save(self, filename):
        """保存文档"""
        self.doc.save(filename)
        print(f"文档已保存: {filename}")


def parse_markdown_to_word(md_file, output_file):
    """解析Markdown文件并生成Word文档"""
    generator = TechnicalReportGenerator()
    
    # 读取Markdown文件
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 解析并构建文档
    lines = content.split('\n')
    i = 0
    in_code_block = False
    in_table = False
    table_headers = []
    table_rows = []
    current_code = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # 处理代码块
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                current_code = []
            else:
                in_code_block = False
                if current_code:
                    generator.add_code_block('\n'.join(current_code))
                current_code = []
            i += 1
            continue
        
        if in_code_block:
            current_code.append(line)
            i += 1
            continue
            
        # 处理表格
        if '|' in line and line.count('|') >= 2:
            if not in_table:
                in_table = True
                table_headers = [cell.strip() for cell in line.split('|') if cell.strip()]
                i += 1
                # 跳过分隔线
                if i < len(lines) and '---' in lines[i]:
                    i += 1
                continue
            else:
                if '---' not in line:
                    table_rows.append([cell.strip() for cell in line.split('|') if cell.strip()])
                else:
                    # 表格结束
                    in_table = False
                    if table_headers and table_rows:
                        generator.add_table(table_headers, table_rows)
                    table_headers = []
                    table_rows = []
                i += 1
                continue
        else:
            if in_table:
                in_table = False
                if table_headers and table_rows:
                    generator.add_table(table_headers, table_rows)
                table_headers = []
                table_rows = []
                
        # 处理标题
        if line.startswith('# '):
            generator.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            generator.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            generator.add_heading(line[4:], level=3)
            
        # 处理分隔线
        elif line.startswith('---'):
            generator.add_page_break()
            
        # 处理列表
        elif line.strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '-', '*', '●')):
            level = (len(line) - len(line.lstrip())) // 4 + 1
            generator.add_list_item(line.strip(), level=level)
            
        # 处理图表
        elif line.startswith('图') and ('：' in line or ':' in line):
            generator.add_figure_placeholder(line)
            
        elif line.startswith('表') and ('：' in line or ':' in line):
            generator.add_figure_placeholder(line)
            
        # 处理空行
        elif not line.strip():
            generator.doc.add_paragraph()
            
        # 处理普通段落
        else:
            if line.strip():
                generator.add_paragraph(line)
                
        i += 1
    
    # 保存文档
    generator.save(output_file)


def main():
    """主函数"""
    md_file = '终端和表计健康状态评估与电量异常识别研究_课题1_技术报告.md'
    output_file = '终端和表计健康状态评估与电量异常识别研究_课题1_技术报告.docx'
    
    if not os.path.exists(md_file):
        print(f"错误: 文件 {md_file} 不存在")
        return
        
    print(f"正在解析 {md_file} ...")
    parse_markdown_to_word(md_file, output_file)
    print("Word文档生成完成！")


if __name__ == '__main__':
    main()
