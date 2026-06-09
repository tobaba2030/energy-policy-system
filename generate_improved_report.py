#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
技术报告Word文档生成器 - 改进版
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def add_heading(doc, text, level=1):
    """添加章节标题"""
    para = doc.add_heading(text, level=level)
    run = para.runs[0] if para.runs else para.add_run(text)
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(18)
    elif level == 2:
        run.font.size = Pt(16)
    elif level == 3:
        run.font.size = Pt(14)


def add_paragraph(doc, text, indent=True):
    """添加正文段落"""
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.5
    if indent:
        para.paragraph_format.first_line_indent = Inches(0.3)
    run = para.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)


def add_list_item(doc, text, level=1):
    """添加列表项"""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3 * level)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)


def add_table(doc, headers, rows):
    """添加表格"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    # 添加表头
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 添加数据行
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = cell_text
            for paragraph in row_cells[i].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = '宋体'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    doc.add_paragraph()


def add_figure_placeholder(doc, caption):
    """添加图表占位符"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"[图: {caption}]")
    run.font.size = Pt(11)
    run.font.italic = True
    doc.add_paragraph()


def main():
    """生成Word文档主函数"""
    doc = Document()
    
    # 设置默认样式
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    font._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    
    # 封面
    add_heading(doc, '终端和表计故障类型与特征关联技术及健康状态评估规则库设计研究', 0)
    doc.add_paragraph()
    add_paragraph(doc, '科技项目技术报告', indent=False)
    doc.add_paragraph()
    
    add_paragraph(doc, '项目名称：终端和表计健康状态评估与电量异常识别研究', indent=False)
    add_paragraph(doc, '课题一名称：终端和表计故障类型与特征关联技术及健康状态评估规则库设计研究', indent=False)
    
    doc.add_page_break()
    
    # 目录
    add_heading(doc, '目录', 1)
    doc.add_paragraph()
    
    add_paragraph(doc, '图目录', indent=False)
    add_list_item(doc, '图1-1 智能量测体系架构示意图', 1)
    add_list_item(doc, '图1-2 课题技术路线图', 1)
    add_list_item(doc, '图2-1 智能终端和表计功能模块分解图', 1)
    add_list_item(doc, '图2-2 故障分类树状图', 1)
    add_list_item(doc, '图2-3 计量芯片漂移故障数学模型示意图', 1)
    add_list_item(doc, '图2-4 故障演化阶段与状态转移图', 1)
    add_list_item(doc, '图2-5 高温高湿环境下电源模块输出纹波变化曲线', 1)
    add_list_item(doc, '图2-6 不同负载率下计量误差的变化趋势图', 1)
    add_list_item(doc, '图3-1 高维特征矩阵构建示意图', 1)
    add_list_item(doc, '图3-2 特征经t-SNE降维后在二维空间的聚类分布图', 1)
    add_list_item(doc, '图3-3 一维卷积神经网络自适应特征学习框架', 1)
    
    doc.add_paragraph()
    
    add_paragraph(doc, '表目录', indent=False)
    add_list_item(doc, '表1-1 国内外研究现状对比表', 1)
    add_list_item(doc, '表2-1 典型故障模式与失效机理对应表', 1)
    add_list_item(doc, '表2-2 多应力下性能退化加速因子表', 1)
    add_list_item(doc, '表3-1 候选特征清单', 1)
    
    doc.add_page_break()
    
    # 摘要
    add_heading(doc, '摘要', 1)
    doc.add_paragraph()
    
    add_paragraph(doc, '随着智能电网建设的深入推进，智能终端和电能表计已成为电力系统的重要组成部分。截至2025年，全国智能电能表安装量已超过6亿只，各类用电信息采集终端数量突破1000万台，构建了全球最大的用电信息采集系统。然而，随着设备规模的快速扩大，终端和表计的健康状态评估面临着严峻挑战：现有评估方法精度不高，告警准确率不足60%；无效告警数量庞大，现场核查工作量巨大；故障机理研究不深入，缺乏统一的评估标准。')
    doc.add_paragraph()
    
    add_paragraph(doc, '本课题针对上述问题，围绕"揭示故障机理-构建特征关联-设计规则库"的研究主线，系统开展终端和表计故障机理与性能退化规律研究，提出多源数据融合的故障特征提取与关联技术，构建科学统一的健康状态评估规则库。主要研究内容包括：（1）基于电力学原理的终端和表计故障机理分析，建立典型故障模式的非线性数学模型；（2）多源状态信息高维特征矩阵构建，基于流形学习的关键特征提取；（3）机理与数据融合的故障类型-特征强关联规则挖掘；（4）全寿命周期分阶段、多尺度健康状态评估规则库设计。')
    doc.add_paragraph()
    
    add_paragraph(doc, '本课题的关键技术创新点在于：提出了机理-数据双驱动的故障特征关联方法，将物理机理分析与数据挖掘技术有机结合；构建了高维特征低维流形提取算法，实现了状态特征的有效降维与可视化；设计了单元级-设备级-台区级三级评估规则库，形成了完整的评估体系。')
    doc.add_paragraph()
    
    add_paragraph(doc, '通过本课题研究，预期可将终端和表计健康状态评估准确率提升至85%以上，无效告警数量压降40%以上，显著提升设备运维效率，保障电力系统安全稳定运行，具有重要的理论意义和工程应用价值。')
    
    doc.add_page_break()
    
    # 第一章
    add_heading(doc, '第一章 绪论', 1)
    doc.add_paragraph()
    
    add_heading(doc, '1.1 项目背景与意义', 2)
    add_heading(doc, '1.1.1 智能量测体系建设与数字电网背景', 3)
    add_paragraph(doc, '在双碳战略目标和新型电力系统建设的背景下，我国智能电网建设取得了举世瞩目的成就。智能量测体系作为智能电网的重要组成部分，经过十余年的建设发展，已形成了覆盖发电、输电、变电、配电、用电各环节的完整量测体系。')
    doc.add_paragraph()
    add_paragraph(doc, '根据国家电网公司2025年统计数据，全国范围内已安装智能电能表超过6亿只，用电信息采集终端数量突破1000万台，采集覆盖率达到99.5%以上，建成了全球规模最大的用电信息采集系统。该系统实现了对电力用户用电信息的实时采集、在线监测和远程控制，为电网公司的电费结算、线损管理、负荷预测、需求响应等业务提供了重要的数据支撑。')
    doc.add_paragraph()
    add_paragraph(doc, '同时，随着数字电网建设的深入推进，智能终端和表计的功能不断拓展，从传统的电能计量向综合感知、边缘计算、双向交互等方向发展。新一代智能电表不仅具备电能计量功能，还集成了电压、电流、功率因数、谐波等多参数量测，支持负荷控制、分布式电源接入管理、电动汽车充电管理等高级应用功能。')
    
    add_figure_placeholder(doc, '智能量测体系架构示意图')
    
    add_heading(doc, '1.1.2 终端与表计健康稳定运行的重大意义', 3)
    add_paragraph(doc, '终端和表计作为电力系统的神经末梢，其健康稳定运行对于保障电力系统安全、提高供电服务质量、提升电网经营效益具有重大意义：')
    doc.add_paragraph()
    add_list_item(doc, '在生产经营方面：智能电能表是电费结算的法定计量器具，其计量准确性直接关系到电力企业和电力用户的经济利益。据统计，因表计故障导致的电量损失每年可达数十亿千瓦时，造成巨大的经济损失。同时，终端设备是用电信息采集系统的关键节点，其稳定运行是确保数据采集完整性和及时性的基础。', 1)
    doc.add_paragraph()
    add_list_item(doc, '在用户利益方面：准确的电能计量是维护电力市场公平公正的前提。终端和表计的健康状态直接影响用电信息的准确性，关系到用户的知情权和公平交易权。通过及时发现和处理设备故障，可以避免因计量不准确引发的用户投诉和纠纷。', 1)
    doc.add_paragraph()
    add_list_item(doc, '在电网安全方面：终端和表计不仅是计量设备，也是电网状态的重要感知节点。通过对终端和表计运行状态的监测，可以及时发现电压异常、负荷过载、谐波超标等电网运行问题，为电网安全稳定运行提供预警信息。同时，终端设备的故障可能导致采集系统瘫痪，影响电网调度决策的准确性。', 1)
    
    add_heading(doc, '1.1.3 当前面临的核心挑战', 3)
    add_paragraph(doc, '尽管智能量测体系建设取得了显著成就，但在终端和表计健康状态评估方面仍面临着诸多挑战：')
    doc.add_paragraph()
    add_list_item(doc, '挑战一：评估精度不高。现有评估方法主要基于单一阈值告警或简单的规则判断，缺乏对设备运行状态的综合分析。据某省级电力公司统计，现有健康状态评估模型的准确率不足60%，存在大量的误判和漏判现象。特别是在设备早期故障阶段，由于故障特征不明显，很难被及时发现。', 1)
    doc.add_paragraph()
    add_list_item(doc, '挑战二：告警数量庞大。随着设备规模的扩大，告警信息呈爆炸式增长。某省级电力公司日均产生终端和表计告警信息超过10万条，其中90%以上为无效告警。运维人员需要花费大量时间进行告警甄别，工作效率低下，真正的故障告警往往被淹没在海量的无效告警中。', 1)
    doc.add_paragraph()
    add_list_item(doc, '挑战三：故障机理不清晰。终端和表计的故障机理复杂，涉及电气、机械、材料、环境等多个因素。目前对典型故障模式的失效机理、演化规律、影响因素等缺乏系统深入的研究，难以建立科学的评估模型和规则。', 1)
    doc.add_paragraph()
    add_list_item(doc, '挑战四：缺乏统一的评估标准。不同地区、不同厂家的评估方法和标准不统一，评估结果缺乏可比性。现有规则库大多基于经验知识，缺乏理论支撑和科学验证，难以保证评估的准确性和一致性。', 1)
    
    # 保存文档
    output_path = 'c:\\AI学习资料\\mesheer\\终端和表计健康状态评估与电量异常识别研究_课题一技术报告.docx'
    doc.save(output_path)
    print(f'Word文档已成功生成：{output_path}')
    return output_path


if __name__ == '__main__':
    main()
