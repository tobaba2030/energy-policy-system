from docx import Document
import sys

def read_docx(file_path):
    try:
        doc = Document(file_path)
        print(f"{'='*80}")
        print(f"文件: {file_path}")
        print(f"{'='*80}\n")
        
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
                full_text.append(para.text)
        
        print(f"\n{'='*80}\n")
        print("表格内容:")
        print(f"{'='*80}\n")
        
        for table_idx, table in enumerate(doc.tables):
            print(f"\n--- 表格 {table_idx + 1} ---\n")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                print(" | ".join(row_data))
            print()
        
        return full_text
    except Exception as e:
        print(f"读取文件错误: {e}")
        return []

if __name__ == "__main__":
    files = [
        r"C:\AI学习资料\科技项目申报\AI+科研项目策划简表生成0517.docx",
        r"C:\AI学习资料\科技项目申报\样例.docx"
    ]
    
    for file in files:
        read_docx(file)
