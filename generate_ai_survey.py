# -*- coding: utf-8 -*-
"""
生成团队AI应用情况交流提纲
"""
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

def set_chinese_font(run, size=12, bold=False):
    """设置中文字体"""
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size)
    run.font.bold = bold

def add_title(doc, text, level=1):
    """添加标题"""
    if level == 1:
        p = doc.add_heading(text, level=level)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_chinese_font(p.runs[0], size=18, bold=True)
    elif level == 2:
        p = doc.add_heading(text, level=level)
        set_chinese_font(p.runs[0], size=16, bold=True)
    elif level == 3:
        p = doc.add_heading(text, level=level)
        set_chinese_font(p.runs[0], size=14, bold=True)
    return p

def add_paragraph(doc, text, indent=False, align_center=False, bold=False, size=12):
    """添加段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_chinese_font(run, size=size, bold=bold)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_list_item(doc, text, indent_level=0, bold=False):
    """添加列表项"""
    p = doc.add_paragraph()
    p.style = 'List Bullet'
    if indent_level > 0:
        p.paragraph_format.left_indent = Inches(indent_level * 0.3)
    run = p.add_run(text)
    set_chinese_font(run, bold=bold)
    return p

def add_table(doc, headers, data, caption=None):
    """添加表格"""
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        set_chinese_font(run, size=11, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 表头
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                set_chinese_font(run, size=11, bold=True)
    
    # 数据行
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
            for para in row_cells[i].paragraphs:
                for run in para.runs:
                    set_chinese_font(run, size=10)
    
    return table

def create_document():
    """创建文档"""
    doc = Document()
    
    # 文档标题页
    add_title(doc, '科创中心团队', level=1)
    add_title(doc, 'AI应用情况交流提纲', level=1)
    add_paragraph(doc, '', align_center=True)
    add_paragraph(doc, '个人数字化AI专项IDP前期调研', align_center=True)
    add_paragraph(doc, '', align_center=True)
    add_paragraph(doc, '编制日期：2026年5月', align_center=True)
    
    doc.add_page_break()
    
    # 一、个人AI应用现状
    add_title(doc, '一、个人AI应用现状', level=1)
    
    add_title(doc, '1.1 AI工具使用情况', level=2)
    add_paragraph(doc, '1. 您目前使用过哪些AI工具？（可多选）', indent=True)
    add_list_item(doc, 'Claude/GPT等大语言模型')
    add_list_item(doc, 'Cursor等AI编程工具')
    add_list_item(doc, 'Gamma等AI生成PPT工具')
    add_list_item(doc, 'Notion等AI知识库工具')
    add_list_item(doc, '其他（请注明）：_________')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '2. 您使用AI工具的频率是？', indent=True)
    add_list_item(doc, '几乎不用')
    add_list_item(doc, '偶尔使用（每周1-2次）')
    add_list_item(doc, '常规使用（每周3-5次）')
    add_list_item(doc, '高频使用（每天都用）')
    
    doc.add_page_break()
    
    # 二、工作场景中的AI应用
    add_title(doc, '二、工作场景中的AI应用', level=1)
    
    add_title(doc, '2.1 业务场景应用', level=2)
    add_paragraph(doc, '3. 在您的日常工作中，AI主要应用在哪些业务场景？', indent=True)
    add_list_item(doc, '申报简表撰写 □')
    add_list_item(doc, '申报指南分析 □')
    add_list_item(doc, '可研报告编写 □')
    add_list_item(doc, '标书文件制作 □')
    add_list_item(doc, '原型设计/Demo开发 □')
    add_list_item(doc, '文献资料整理 □')
    add_list_item(doc, '其他（请注明）：_________')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '4. 请分享一个您使用AI工具提升工作效率的具体案例。', indent=True, bold=True)
    
    # 添加空白区域供填写
    for i in range(5):
        add_paragraph(doc, '')
    
    doc.add_page_break()
    
    # 三、AI应用效果评估
    add_title(doc, '三、AI应用效果评估', level=1)
    
    add_title(doc, '3.1 效率提升评估', level=2)
    add_paragraph(doc, '5. 您认为AI工具对您工作效率的提升程度是？', indent=True)
    add_list_item(doc, '不明显（0-10%）')
    add_list_item(doc, '有一定提升（10-30%）')
    add_list_item(doc, '显著提升（30-50%）')
    add_list_item(doc, '大幅提升（50%以上）')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '6. 在使用AI工具过程中，您遇到过哪些挑战？', indent=True)
    add_list_item(doc, '工具操作难度大')
    add_list_item(doc, '生成内容质量不稳定')
    add_list_item(doc, '缺乏行业专业知识')
    add_list_item(doc, '数据安全顾虑')
    add_list_item(doc, '其他（请注明）：_________')
    
    doc.add_page_break()
    
    # 四、技能提升与未来规划
    add_title(doc, '四、技能提升与未来规划', level=1)
    
    add_title(doc, '4.1 学习需求', level=2)
    add_paragraph(doc, '7. 您目前最想学习/提升的AI技能是？', indent=True)
    add_list_item(doc, '提示词工程')
    add_list_item(doc, 'AI辅助编程')
    add_list_item(doc, '数据分析可视化')
    add_list_item(doc, '自动化工作流')
    add_list_item(doc, '其他（请注明）：_________')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '8. 您对团队AI能力建设有哪些建议？', indent=True, bold=True)
    for i in range(4):
        add_paragraph(doc, '')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '9. 您希望获得哪些资源或支持来提升AI应用能力？', indent=True, bold=True)
    for i in range(4):
        add_paragraph(doc, '')
    
    doc.add_page_break()
    
    # 五、Demo开发专项（针对负责Demo开发的成员）
    add_title(doc, '五、Demo开发专项', level=1)
    add_paragraph(doc, '（仅针对负责Demo开发的成员：周杰、吴敏、孙浩）', align_center=True, size=10)
    
    add_paragraph(doc, '')
    add_paragraph(doc, '10. 您在科技项目Demo开发中，AI工具发挥了哪些作用？', indent=True, bold=True)
    for i in range(4):
        add_paragraph(doc, '')
    
    add_paragraph(doc, '')
    add_paragraph(doc, '11. 请分享一个AI辅助Demo开发的成功案例。', indent=True, bold=True)
    for i in range(4):
        add_paragraph(doc, '')
    
    doc.add_page_break()
    
    # 六、交流记录表
    add_title(doc, '六、交流记录表', level=1)
    
    record_headers = ['姓名', '交流时间', '记录人', '备注']
    record_data = [
        ['张伟', '', '', ''],
        ['李明', '', '', ''],
        ['王芳', '', '', ''],
        ['刘洋', '', '', ''],
        ['陈静', '', '', ''],
        ['赵强', '', '', ''],
        ['周杰', '', '', ''],
        ['吴敏', '', '', ''],
        ['孙浩', '', '', ''],
        ['朱婷', '', '', '']
    ]
    add_table(doc, record_headers, record_data, '表1 团队成员交流记录')
    
    add_paragraph(doc, '')
    add_title(doc, '交流方式建议', level=2)
    add_list_item(doc, '形式：可以采用一对一访谈或小组座谈会的形式')
    add_list_item(doc, '时间：每位成员建议交流时间15-20分钟')
    add_list_item(doc, '记录：安排专人记录，后续整理成《团队AI应用现状分析报告》')
    add_list_item(doc, '跟进：根据交流结果，制定个性化的AI能力提升计划')
    
    doc.add_page_break()
    
    # 文档信息页
    add_paragraph(doc, '============================================================', align_center=True)
    add_paragraph(doc, '文档信息', align_center=True)
    add_paragraph(doc, '文档名称：科创中心团队AI应用情况交流提纲', align_center=True)
    add_paragraph(doc, '编制日期：2026年5月', align_center=True)
    add_paragraph(doc, '用途：个人数字化AI专项IDP前期调研', align_center=True)
    add_paragraph(doc, '============================================================', align_center=True)
    
    return doc

if __name__ == '__main__':
    print('正在生成Word文档...')
    doc = create_document()
    output_path = 'c:/AI学习资料/mesheer/科创中心团队AI应用情况交流提纲.docx'
    doc.save(output_path)
    print(f'文档生成成功！保存位置：{output_path}')
