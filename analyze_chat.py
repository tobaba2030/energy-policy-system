import openpyxl

file_path = r'C:\Users\jianlinw\Desktop\手机助理导出的信息2026-05-04_135059138.xlsx'
wb = openpyxl.load_workbook(file_path)
print('Sheet names:', wb.sheetnames)

ws = wb.active
print('Active sheet:', ws.title)
print('Max row:', ws.max_row)
print('Max col:', ws.max_column)
print('\nFirst 10 rows:')
for i, row in enumerate(ws.iter_rows(max_row=10, values_only=True), 1):
    print(f'Row {i}:', row)
