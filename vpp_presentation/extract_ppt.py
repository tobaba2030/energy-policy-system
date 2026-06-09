
from pptx import Presentation
import os
from zipfile import ZipFile
import shutil

# 读取原PPT
ppt_path = r"C:\Users\jianlinw\Desktop\南方电网《虚拟电厂技术的探索与实践》 (1).pptx"
prs = Presentation(ppt_path)

print("幻灯片数量:", len(prs.slides))
print("幻灯片宽高:", prs.slide_width, prs.slide_height)
print("=" * 80)

# 提取所有图片
output_dir = r"c:\AI学习资料\mesheer\vpp_presentation\extracted_images"
if not os.path.exists(output_dir):
    os.makedirs(output_dir)

# 从PPTX中提取媒体文件
with ZipFile(ppt_path, 'r') as zip_ref:
    media_files = [f for f in zip_ref.namelist() if f.startswith('ppt/media/')]
    print("找到", len(media_files), "个媒体文件")
    
    for i, media_file in enumerate(media_files, 1):
        ext = os.path.splitext(media_file)[1]
        dest_path = os.path.join(output_dir, "image_" + str(i) + ext)
        with zip_ref.open(media_file) as source, open(dest_path, 'wb') as target:
            shutil.copyfileobj(source, target)
        print("  提取图片:", media_file, "-&gt;", dest_path)

print("=" * 80)

# 分析每个幻灯片
for slide_idx, slide in enumerate(prs.slides, 1):
    print("\n=== 幻灯片", slide_idx, "===")
    
    # 统计形状类型
    shape_types = {}
    for shape in slide.shapes:
        stype = str(type(shape).__name__)
        if stype in shape_types:
            shape_types[stype] = shape_types[stype] + 1
        else:
            shape_types[stype] = 1
    
    print("  形状类型统计:", shape_types)
    
    # 检查图片和文本
    for shape_idx, shape in enumerate(slide.shapes, 1):
        if hasattr(shape, 'image'):
            print("  形状", shape_idx, ": 图片")
        elif shape.has_text_frame:
            text = shape.text.strip()
            if text:
                preview = text[:80]
                print("  形状", shape_idx, ": 文本 -", preview)
