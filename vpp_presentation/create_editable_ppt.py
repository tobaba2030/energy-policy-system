
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

# 输入路径
original_ppt_path = r"C:\Users\jianlinw\Desktop\南方电网《虚拟电厂技术的探索与实践》 (1).pptx"
images_dir = r"c:\AI学习资料\mesheer\vpp_presentation\extracted_images"

# 输出路径
output_ppt_path = r"c:\AI学习资料\mesheer\vpp_presentation\虚拟电厂_可编辑版.pptx"

# 读取原PPT获取尺寸
original_prs = Presentation(original_ppt_path)
prs = Presentation()

# 设置幻灯片尺寸与原PPT一致
prs.slide_width = original_prs.slide_width
prs.slide_height = original_prs.slide_height

# 转换EMU到英寸
width_in = prs.slide_width / 914400
height_in = prs.slide_height / 914400

print("创建可编辑PPT...")
print(f"幻灯片尺寸: {width_in:.2f} x {height_in:.2f} 英寸")

# 处理每张幻灯片
for slide_num in range(1, 32):
    print(f"处理幻灯片 {slide_num}/31...")
    
    # 使用空白布局
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # 1. 添加原始图片作为背景（全屏）
    img_path = os.path.join(images_dir, f"image_{slide_num}.jpg")
    if os.path.exists(img_path):
        slide.shapes.add_picture(
            img_path,
            left=0,
            top=0,
            width=prs.slide_width,
            height=prs.slide_height
        )
    
    # 2. 添加可编辑的文本占位区域
    # 添加标题文本框
    title_box = slide.shapes.add_textbox(
        left=Inches(0.5), 
        top=Inches(0.3), 
        width=prs.slide_width - Inches(1), 
        height=Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.text = "【在此输入标题】"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER
    title_box.fill.background()  # 透明背景
    
    # 添加内容文本框
    content_box = slide.shapes.add_textbox(
        left=Inches(0.7), 
        top=Inches(1.2), 
        width=prs.slide_width - Inches(1.4), 
        height=prs.slide_height - Inches(2)
    )
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    content_para = content_frame.paragraphs[0]
    content_para.text = "【在此输入内容】\n\n提示：\n1. 此文本框可编辑\n2. 可以添加多行内容\n3. 调整字体大小和格式"
    content_para.font.size = Pt(18)
    content_box.fill.background()  # 透明背景
    
    # 添加小提示标签
    note_box = slide.shapes.add_textbox(
        left=Inches(0.2), 
        top=Inches(height_in - 0.4), 
        width=Inches(3), 
        height=Inches(0.3)
    )
    note_frame = note_box.text_frame
    note_frame.text = f"幻灯片 {slide_num} - 可编辑版本"
    note_para = note_frame.paragraphs[0]
    note_para.font.size = Pt(10)
    # 使用默认颜色，无需设置

# 保存PPT
prs.save(output_ppt_path)
print(f"\n成功！可编辑PPT已保存至: {output_ppt_path}")
print("\n说明：")
print("- 每张幻灯片保留了原始图片作为背景")
print("- 添加了标题和内容两个可编辑文本框")
print("- 您可以在PowerPoint中打开并编辑这些文本框")
