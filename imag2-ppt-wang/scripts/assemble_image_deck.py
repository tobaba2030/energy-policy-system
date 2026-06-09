#!/usr/bin/env python3
"""Assemble image-only body slides into a reference PPTX.

The first slide and final slide are preserved by default. Each body slide is
cleared and replaced by exactly one full-slide image. This is intentionally
strict so the final deck remains image2-driven and auditable.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from pptx.dml.color import RGBColor
from pptx import Presentation


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def body_image_paths(images_dir: Path) -> list[Path]:
    files = [p for p in images_dir.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXTS]
    return sorted(files, key=lambda p: p.name)


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)


def set_text(shape, value: str, *, force_white_bold: bool = False) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = value
    if force_white_bold:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)


def replace_cover_title(prs: Presentation, cover_title: str | None) -> str:
    if not cover_title:
        return "not_requested"
    slide = prs.slides[0]
    candidates = []
    blocked = ("演亦文化", "恒心", "打造全国领先", "谢谢", "有限责任公司")
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = " ".join(shape.text.split())
        if not text:
            continue
        if any(token in text for token in blocked):
            continue
        area = int(shape.width) * int(shape.height)
        candidates.append((area, shape, text))
    if not candidates:
        return "not_found"
    _, shape, old_text = max(candidates, key=lambda item: item[0])
    set_text(shape, cover_title, force_white_bold=True)
    return f"replaced_white_bold: {old_text} -> {cover_title}"


def slide_texts(slide) -> list[str]:
    texts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            texts.append(" ".join(shape.text.split()))
    return texts


def write_audit(
    audit_path: Path,
    reference_pptx: Path,
    out_pptx: Path,
    images: Iterable[Path],
    prs: Presentation,
    cover_status: str,
) -> None:
    lines = [
        "# Image2 PPT Assembly Audit",
        "",
        "## Inputs",
        f"- reference_pptx: `{reference_pptx}`",
        f"- out_pptx: `{out_pptx}`",
        f"- image_count: {len(list(images))}",
        f"- slide_count: {len(prs.slides)}",
        f"- cover_title_status: {cover_status}",
        "",
        "## Body Slide Policy",
        "- Slide 1 preserved as cover.",
        "- Final slide preserved as closing slide.",
        "- Every body slide is cleared and replaced by exactly one full-slide image.",
        "- No PPT-native body text, charts, labels, or overlays are added.",
        "",
        "## Images",
    ]
    for idx, image in enumerate(images, start=2):
        lines.append(f"- page {idx:02d}: `{image}`")
    lines.extend(["", "## Final Slide Text Samples"])
    for idx, slide in enumerate(prs.slides, start=1):
        texts = slide_texts(slide)
        sample = " | ".join(texts[:4]) if texts else "(image-only or no text)"
        lines.append(f"- slide {idx:02d}: {sample}")
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference-pptx", required=True, type=Path)
    parser.add_argument("--images-dir", required=True, type=Path)
    parser.add_argument("--out-pptx", required=True, type=Path)
    parser.add_argument("--audit-markdown", type=Path)
    parser.add_argument("--cover-title")
    args = parser.parse_args()

    prs = Presentation(str(args.reference_pptx))
    if len(prs.slides) < 3:
        raise SystemExit("Reference PPTX must contain at least cover, one body slide, and closing slide.")

    images = body_image_paths(args.images_dir)
    body_count = len(prs.slides) - 2
    if len(images) != body_count:
        raise SystemExit(f"Image count mismatch: got {len(images)}, expected {body_count} body slides.")

    cover_status = replace_cover_title(prs, args.cover_title)

    for offset, image in enumerate(images, start=1):
        slide = prs.slides[offset]
        clear_slide(slide)
        slide.shapes.add_picture(str(image), 0, 0, prs.slide_width, prs.slide_height)

    args.out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out_pptx))

    if args.audit_markdown:
        write_audit(args.audit_markdown, args.reference_pptx, args.out_pptx, images, prs, cover_status)

    print(args.out_pptx)


if __name__ == "__main__":
    main()
