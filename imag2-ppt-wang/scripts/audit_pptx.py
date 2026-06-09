#!/usr/bin/env python3
"""Write a concise Markdown inventory for a PPTX."""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

from pptx import Presentation


def shape_type_name(shape) -> str:
    return str(shape.shape_type).replace("\n", " ")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    prs = Presentation(str(args.pptx))
    lines = [
        "# PPTX Audit",
        "",
        f"- pptx: `{args.pptx}`",
        f"- slides: {len(prs.slides)}",
        f"- size_emu: {prs.slide_width} x {prs.slide_height}",
        "",
        "## Slides",
    ]

    for idx, slide in enumerate(prs.slides, start=1):
        counts = Counter(shape_type_name(shape) for shape in slide.shapes)
        lines.append(f"### Slide {idx:02d}")
        lines.append(f"- shape_count: {len(slide.shapes)}")
        lines.append(f"- shape_types: {dict(counts)}")
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                texts.append(" ".join(shape.text.split()))
        if texts:
            lines.append("- text_samples:")
            for text in texts[:10]:
                lines.append(f"  - {text[:240]}")
        else:
            lines.append("- text_samples: []")
        lines.append("")

    content = "\n".join(lines)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(content + "\n", encoding="utf-8")
        print(args.out)
    else:
        print(content)


if __name__ == "__main__":
    main()

